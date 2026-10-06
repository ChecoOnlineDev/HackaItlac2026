"""Ensayo con un Excel ajeno (docs/recursos/comprasejer2026.xlsx, 350 filas) en modo ALTA."""

import warnings
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.modulos.importacion.categorias_sugeridas import es_servicio
from tests.importacion.ayudas import IMPORTACION, VISTA_PREVIA, articulo, confirmacion, conteos

RUTA = Path(__file__).resolve().parents[3] / "docs" / "recursos" / "comprasejer2026.xlsx"
COLUMNAS = {"nombre": 1, "cantidad": 2, "costo": 3}  # la columna 0 es una clave de producto


def _filas() -> list[list[str]]:
    if not RUTA.exists():
        pytest.skip("No está el Excel de ejemplo")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        libro = load_workbook(RUTA, read_only=True, data_only=True)
    filas = list(libro.active.iter_rows(values_only=True))[1:]
    return [["" if c is None else str(c) for c in f[:4]] for f in filas]


def test_I_14_las_350_filas_del_excel_real_pasan_por_el_analizador_en_alta(compras, session):
    filas = _filas()
    assert len(filas) == 350
    antes = conteos(session)
    r = compras.post(
        VISTA_PREVIA,
        json={"modo": "ALTA", "filas": filas, "columnas": COLUMNAS, "primera_fila": 2},
    )
    assert r.status_code == 200, r.text
    vp = r.json()
    assert conteos(session) == antes  # la vista previa no escribe
    res = vp["resumen"]
    assert res["total"] == 350
    assert res["validas"] + res["con_error"] + res["excluidas"] + res["vacias"] == 350 - sum(
        len(f["unida_de"]) for f in vp["filas_validas"]
    )
    # Los servicios se excluyen, no son error.
    esperadas = [i + 2 for i, f in enumerate(filas) if es_servicio(f[1])]
    assert esperadas and [x["fila"] for x in vp["filas_excluidas"]] == esperadas
    # Cobertura de categorías: toda fila trae sugerencia o queda «por revisar».
    unidas = {n for f in vp["filas_validas"] for n in f["unida_de"]}
    sin_cubrir = []
    for f in vp["filas_validas"]:
        assert f["categoria"] is None and f["categoria_sugerida"] is not None
    for f in vp["filas_error"]:
        codigos = {m["codigo"] for m in f["motivos"]}
        if f["fila"] in unidas or f.get("categoria_sugerida") or "CATEGORIA_DESCONOCIDA" in codigos:
            continue
        sin_cubrir.append(f["fila"])
    assert sin_cubrir == []
    assert res["por_revisar"] == len(
        [f for f in vp["filas_error"] if "categoria_sugerida" not in f]
    )


def test_I_14_aceptadas_las_sugerencias_el_excel_real_se_importa(compras, session):
    filas = _filas()
    cuerpo = {"modo": "ALTA", "filas": filas, "columnas": COLUMNAS, "primera_fila": 2}
    vp = compras.post(VISTA_PREVIA, json=cuerpo).json()
    aceptadas = {
        str(f["fila"]): f["categoria_sugerida"]["id"]
        for f in vp["filas_validas"]
        if f["categoria_sugerida"]
    }
    r = compras.post(
        IMPORTACION,
        json=confirmacion(filas, columnas=COLUMNAS, primera_fila=2, categoria_por_fila=aceptadas),
    )
    assert r.status_code == 201, r.text
    salida = r.json()
    assert salida["resumen"]["articulos_creados"] > 100
    assert all(a["codigo_generado"] for a in salida["articulos_creados"])
    assert articulo(session, salida["articulos_creados"][0]["codigo"]) is not None
    # Con las categorías elegidas, lo que era sugerencia pendiente ahora vuelve con categoría.
    otra = compras.post(VISTA_PREVIA, json={**cuerpo, "categoria_por_fila": aceptadas}).json()
    assert all(f["categoria"] is not None for f in otra["filas_validas"] if f["fila"] in aceptadas)
