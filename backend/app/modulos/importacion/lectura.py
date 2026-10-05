"""Lectura de la tabla: texto de las celdas, mapeo de columnas y archivos `.xlsx`.

Nada aqui toca la base de datos. El `.xlsx` se lee en memoria con `openpyxl` en modo de solo
lectura y con los valores guardados (`data_only=True`): una formula nunca se ejecuta, se lee el
ultimo resultado que Excel guardo (o queda vacia). Se rechazan los libros con macros, los que no
son un `.xlsx` valido por su contenido, los muy grandes y los que se descomprimen demasiado
(zip bomb). Nunca se escribe nada en disco.
"""

import io
import re
import unicodedata
import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from app.modulos.importacion.exceptions import ArchivoInvalido
from app.modulos.importacion.schemas import CAMPOS, MAX_CELDA, MAX_COLUMNAS, MAX_FILAS

MAX_BYTES_ARCHIVO = 5 * 1024 * 1024
MAX_BYTES_DESCOMPRIMIDO = 50 * 1024 * 1024
MAX_ENTRADAS_ZIP = 1000
# Filas que se recorren (vacias incluidas) antes de rendirse: una hoja con la celda XFD1048576
# marcada no debe tardar minutos.
MAX_FILAS_RECORRIDAS = MAX_FILAS * 4 + 1

_FIRMA_ZIP = b"PK\x03\x04"
_CONTENT_TYPE_LIBRO = "spreadsheetml.sheet.main+xml"
_CONTENT_TYPE_MACROS = "macroEnabled"


# ------------------------------------------------------------------------------ texto


def clave(texto: str) -> str:
    """Forma de comparar nombres: sin acentos, sin mayusculas y con espacios simples. La base
    compara igual (`utf8mb4_0900_ai_ci`), asi que un codigo `Taladro` y `TALADRO` son el mismo."""
    descompuesto = unicodedata.normalize("NFKD", texto)
    sin_acentos = "".join(c for c in descompuesto if not unicodedata.combining(c))
    return " ".join(sin_acentos.casefold().split())


def texto_de_celda(valor: Any) -> str:
    """El texto de una celda: los numeros enteros sin `.0` (el Excel guarda `1001` como
    `1001.0`), las fechas en ISO, los vacios como `""`. Quita espacios y caracteres de control."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if valor != valor or valor in (float("inf"), float("-inf")):
            return ""
        if valor == int(valor) and abs(valor) < 1e15:
            return str(int(valor))
        return format(Decimal(repr(valor)), "f")
    if isinstance(valor, datetime | date | time):
        return valor.isoformat()
    texto = str(valor)
    texto = "".join(
        c for c in texto if c in "\t\n\r" or not unicodedata.category(c).startswith("C")
    )
    return " ".join(texto.split())


# ----------------------------------------------------------------- mapeo de columnas

# Palabras (ya sin acentos) que delatan que dato trae una columna. El orden importa: el dato
# mas especifico primero (`codigo de pieza` antes que `codigo`).
_PALABRAS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("serie", ("serie", "serial")),
    ("costo", ("costo", "precio", "importe")),
    ("cantidad", ("cantidad", "cant", "existencia", "existencias", "stock", "unidades")),
    ("almacen", ("almacen", "bodega")),
    ("categoria", ("categoria", "familia", "rubro")),
    ("marca", ("marca", "fabricante")),
    ("nombre", ("nombre", "descripcion", "producto", "material", "herramienta")),
)
_CODIGO = ("codigo", "clave", "sku", "folio", "etiqueta", "qr")


def proponer_columnas(encabezados: list[str | None]) -> dict[str, int | None]:
    """Relaciona cada dato con la columna cuyo encabezado se le parece (US-IMP-001). Una
    columna sirve a un solo dato; los datos que no se reconocen quedan en `None`."""
    columnas: dict[str, int | None] = dict.fromkeys(CAMPOS)
    usadas: set[int] = set()
    palabras = [set(re.findall(r"[a-z0-9]+", clave(e or ""))) for e in encabezados]

    def tomar(campo: str, buscadas: tuple[str, ...], *, exigir: tuple[str, ...] = ()) -> None:
        for i, conjunto in enumerate(palabras):
            if i in usadas or columnas[campo] is not None:
                continue
            if conjunto & set(buscadas) and (not exigir or conjunto & set(exigir)):
                columnas[campo] = i
                usadas.add(i)

    # `codigo pieza`, `etiqueta pieza`, `id pieza`: palabra de codigo mas "pieza".
    tomar("codigo_pieza", (*_CODIGO, "id"), exigir=("pieza", "piezas"))
    for campo, buscadas in _PALABRAS:
        tomar(campo, buscadas)
    tomar("codigo", _CODIGO)
    return columnas


# ----------------------------------------------------------------------------- xlsx


def _invalido(mensaje: str) -> ArchivoInvalido:
    return ArchivoInvalido(mensaje, {"campo": "archivo"})


def _revisar_zip(contenido: bytes) -> None:
    """Revisa el contenedor antes de abrirlo con `openpyxl`: es un zip de verdad, es un libro de
    Excel sin macros y no se descomprime a un tamano absurdo."""
    if not contenido.startswith(_FIRMA_ZIP):
        raise _invalido("El archivo no es un Excel (.xlsx) válido.")
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ENTRADAS_ZIP:
                raise _invalido("El archivo tiene demasiadas partes internas.")
            if sum(i.file_size for i in infos) > MAX_BYTES_DESCOMPRIMIDO:
                raise _invalido("El archivo es demasiado grande al abrirlo.")
            if any(i.flag_bits & 0x1 for i in infos):
                raise _invalido("El archivo está protegido con contraseña.")
            nombres = {i.filename.lower() for i in infos}
            if "[content_types].xml" not in nombres or "xl/workbook.xml" not in nombres:
                raise _invalido("El archivo no es un Excel (.xlsx) válido.")
            if any(n.endswith("vbaproject.bin") or "/vba" in n for n in nombres):
                raise _invalido("El archivo trae macros. Guárdalo como .xlsx sin macros.")
            tipos = zf.read("[Content_Types].xml")
            if len(tipos) > 1_000_000:
                raise _invalido("El archivo no es un Excel (.xlsx) válido.")
            if _CONTENT_TYPE_MACROS.encode() in tipos:
                raise _invalido("El archivo trae macros. Guárdalo como .xlsx sin macros.")
            if _CONTENT_TYPE_LIBRO.encode() not in tipos:
                raise _invalido("El archivo no es un Excel (.xlsx) válido.")
    except zipfile.BadZipFile as exc:
        raise _invalido("El archivo no es un Excel (.xlsx) válido.") from exc


def leer_xlsx(nombre_archivo: str | None, contenido: bytes) -> tuple[str | None, list[list[str]]]:
    """Convierte un `.xlsx` en `(nombre de la hoja, filas de texto)`; la primera hoja. Las filas
    vacias del final se quitan; las vacias de en medio se conservan (conservan su numero de fila).
    Cualquier problema es `ArchivoInvalido` con un mensaje en espanol."""
    nombre = (nombre_archivo or "").strip().lower()
    if nombre.endswith(".xlsm") or nombre.endswith(".xlsb") or nombre.endswith(".xls"):
        raise _invalido("Solo se admite Excel con extensión .xlsx (sin macros).")
    if not nombre.endswith(".xlsx"):
        raise _invalido("Sube un archivo de Excel con extensión .xlsx.")
    if not contenido:
        raise _invalido("El archivo está vacío.")
    if len(contenido) > MAX_BYTES_ARCHIVO:
        raise _invalido(f"El archivo pesa más de {MAX_BYTES_ARCHIVO // (1024 * 1024)} MB.")
    _revisar_zip(contenido)

    from openpyxl import load_workbook

    libro = None
    try:
        libro = load_workbook(
            io.BytesIO(contenido), read_only=True, data_only=True, keep_links=False
        )
        hoja = libro.active
        if hoja is None:
            return None, []
        filas: list[list[str]] = []
        recorridas = 0
        for crudas in hoja.iter_rows(values_only=True):
            recorridas += 1
            if recorridas > MAX_FILAS_RECORRIDAS:
                raise _invalido(f"La hoja tiene demasiadas filas (máximo {MAX_FILAS}).")
            celdas = [texto_de_celda(c) for c in crudas]
            while celdas and not celdas[-1]:
                celdas.pop()
            if len(celdas) > MAX_COLUMNAS:
                raise _invalido(f"La hoja tiene más de {MAX_COLUMNAS} columnas.")
            if any(len(c) > MAX_CELDA for c in celdas):
                raise _invalido(f"Una celda pasa de {MAX_CELDA} caracteres.")
            filas.append(celdas)
        while filas and not filas[-1]:
            filas.pop()
        if len(filas) > MAX_FILAS + 1:  # el encabezado no cuenta
            raise _invalido(f"La hoja tiene demasiadas filas (máximo {MAX_FILAS}).")
        return hoja.title, filas
    except ArchivoInvalido:
        raise
    except Exception as exc:  # noqa: BLE001 - un archivo malo nunca debe ser un 500
        raise _invalido("No se pudo leer el archivo. Verifica que sea un Excel (.xlsx).") from exc
    finally:
        if libro is not None:
            libro.close()
